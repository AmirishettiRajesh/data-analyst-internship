import os
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from docx import Document

os.makedirs('outputs/charts', exist_ok=True)
os.makedirs('report', exist_ok=True)

df = pd.read_csv('dataset.csv')

summary_stats = df.describe()
dept_avg = df.groupby('Department')[['Salary', 'Performance_Score']].mean()

with open('outputs/results.txt', 'w') as f:
    f.write('=== DATASET SUMMARY STATISTICS ===\n')
    f.write(summary_stats.to_string())
    f.write('\n\n=== AVERAGE METRICS BY DEPARTMENT ===\n')
    f.write(dept_avg.to_string())

plt.style.use('seaborn-v0_8-whitegrid' if 'seaborn-v0_8-whitegrid' in plt.style.available else 'default')

plt.figure(figsize=(8, 5))
sns.histplot(df['Salary'], kde=True, color='skyblue')
plt.title('Salary Distribution')
plt.xlabel('Salary ($)')
plt.ylabel('Frequency')
plt.tight_layout()
plt.savefig('outputs/charts/histogram.png', dpi=300)
plt.close()

plt.figure(figsize=(8, 5))
dept_avg['Salary'].plot(kind='bar', color='teal')
plt.title('Average Salary by Department')
plt.xlabel('Department')
plt.ylabel('Average Salary ($)')
plt.xticks(rotation=0)
plt.tight_layout()
plt.savefig('outputs/charts/bar_chart.png', dpi=300)
plt.close()

plt.figure(figsize=(8, 5))
sns.scatterplot(data=df, x='Experience_Years', y='Salary', hue='Department', s=100)
plt.title('Experience vs. Salary')
plt.xlabel('Years of Experience')
plt.ylabel('Salary ($)')
plt.tight_layout()
plt.savefig('outputs/charts/scatter_plot.png', dpi=300)
plt.close()

plt.figure(figsize=(8, 5))
sns.boxplot(data=df, x='Sales_Target_Met', y='Performance_Score', palette='Set2')
plt.title('Performance Score Distribution by Sales Target Status')
plt.xlabel('Sales Target Met')
plt.ylabel('Performance Score')
plt.tight_layout()
plt.savefig('outputs/charts/box_plot.png', dpi=300)
plt.close()

doc = Document()
doc.add_heading('Task 2: Hands-On Data Lab Report', level=0)

doc.add_heading('1. Executive Summary', level=1)
doc.add_paragraph('This report documents the implementation of the Hands-On Data Lab environment.')

doc.add_heading('2. Statistical Analysis Results', level=1)
doc.add_paragraph(f'Total Employees Analyzed: {len(df)}')
doc.add_paragraph(f'Average Salary: ')
doc.add_paragraph(f'Average Performance Score: {df["Performance_Score"].mean():.2f}')

doc.add_heading('3. Key Visualizations', level=1)
for chart_name in ['histogram.png', 'bar_chart.png', 'scatter_plot.png', 'box_plot.png']:
    chart_path = os.path.join('outputs/charts', chart_name)
    if os.path.exists(chart_path):
        doc.add_paragraph(f'Chart: {chart_name}')
        doc.add_picture(chart_path)

doc.save('report/Task_2_Hands_On_Data_Lab_Report.docx')
print('Task 2 execution completed successfully!')
